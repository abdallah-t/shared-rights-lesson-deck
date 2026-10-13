"""Builds the picture PPTX from the images made by scripts/capture-steps.ts.

  uv run --no-project --with python-pptx python scripts/build-pptx.py <steps-dir> <out.pptx>

One 16:9 slide per image, full bleed. A new deck slide pushes in from the
right; a step within the same slide fades, so reveals feel like the original.
"""
import json, pathlib, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

steps_dir, out = pathlib.Path(sys.argv[1]), sys.argv[2]
steps = json.loads((steps_dir / "steps.json").read_text())
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
prev = None
for s in steps:
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(steps_dir / s["file"]), 0, 0, prs.slide_width, prs.slide_height)
    kind = '<p:fade/>' if s["slide"] == prev else '<p:push dir="l"/>'
    speed = "fast" if s["slide"] == prev else "med"
    tr = etree.fromstring(f'<p:transition xmlns:p="{P}" spd="{speed}">{kind}</p:transition>')
    slide._element.find(f"{{{P}}}clrMapOvr").addnext(tr)  # schema order: after clrMapOvr, before timing
    prev = s["slide"]
prs.core_properties.title = "الحقوق المشتركة بين الزوجين"
prs.save(out)
print(len(steps), "slides ->", out)
