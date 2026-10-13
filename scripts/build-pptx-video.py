"""Builds the animated PPTX (v2) from the clips made by scripts/capture-clips.ts.

  uv run --no-project --with python-pptx python scripts/build-pptx-video.py <clips-dir> <out.pptx>

One 16:9 slide per step. Each slide is a full-bleed video of that step's real
animation, set to play by itself as the slide opens and to rest on its last
frame. Slides cut straight from one to the next: the red wipe is in the video.
"""
import json, pathlib, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

clips_dir, out = pathlib.Path(sys.argv[1]), sys.argv[2]
clips = json.loads((clips_dir / "clips.json").read_text())
P = "http://schemas.openxmlformats.org/presentationml/2006/main"

# What PowerPoint itself writes for a video set to "Start: Automatically".
TIMING = """<p:timing xmlns:p="%(P)s"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="%(ms)d" fill="hold"/><p:tgtEl><p:spTgt spid="%(id)d"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="0" showWhenStopped="1"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="%(id)d"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]
for c in clips:
    slide = prs.slides.add_slide(blank)
    mp4 = clips_dir / c["file"]
    movie = slide.shapes.add_movie(str(mp4), 0, 0, prs.slide_width, prs.slide_height, poster_frame_image=str(mp4.with_suffix(".jpg")), mime_type="video/mp4")
    old = slide._element.find(f"{{{P}}}timing")
    if old is not None:
        slide._element.remove(old)
    slide._element.append(etree.fromstring(TIMING % {"P": P, "ms": round(c["seconds"] * 1000), "id": movie.shape_id}))
prs.core_properties.title = "الحقوق المشتركة بين الزوجين"
prs.save(out)
print(len(clips), "slides ->", out)
